#!/usr/bin/env python3
"""Biçim testleri için örnek dosyaları üretir.

Kaynak çalışma kitapları openpyxl ile yazılır, LibreOffice (soffice) ile
.ods, .ots, .xls, .xltx ve .xlsm biçimlerine çevrilir. Çalıştırma:

    python3 apps/tablo/src/io/__fixtures__/uret.py

Gerekenler: python3 + openpyxl, LibreOffice (soffice). Üretilen dosyalar bu
dizine yazılır; ara .xlsx dosyaları geçici dizinde kalır.
"""
import datetime
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))


def ornek_kitap(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Satışlar"
    ince = Side(style="thin", color="000000")
    ws.append(["Ürün", "Adet", "Fiyat", "Tutar", "Tarih", "Oran", "Not"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center")
        c.border = Border(bottom=ince)
    rows = [
        ("Çay", 10, 2.5, datetime.date(2024, 3, 15), 0.25, "İzmir"),
        ("Şeker", 4, 30, datetime.date(2024, 1, 2), 0.125, "Muğla"),
        ("Kahve", 7, 12.75, datetime.date(2023, 12, 31), 0.5, "ığüşöç ĞÜŞİÖÇ"),
    ]
    for i, (ad, adet, fiyat, tarih, oran, notu) in enumerate(rows, start=2):
        ws.cell(i, 1, ad)
        ws.cell(i, 2, adet)
        ws.cell(i, 3, fiyat).number_format = "#,##0.00"
        ws.cell(i, 4, f"=B{i}*C{i}").number_format = "#,##0.00"
        ws.cell(i, 5, tarih).number_format = "DD.MM.YYYY"
        ws.cell(i, 6, oran).number_format = "0.00%"
        ws.cell(i, 7, notu)
    ws["A5"] = "Toplam"
    ws["A5"].font = Font(italic=True, underline="single", size=14, color="C00000")
    ws["D5"] = "=SUM(D2:D4)"
    ws["D5"].border = Border(top=ince, bottom=Side(style="double", color="000000"))
    ws["B6"] = "=AVERAGE(B2:B4)"
    ws["C6"] = '=IF(D5>100,"Büyük","Küçük")'
    ws["D6"] = "=ROUND(C4*1.18,2)"
    ws["E6"] = "=VLOOKUP(\"Şeker\",A2:C4,3,FALSE)"
    ws["F6"] = '=CONCATENATE(A2," ve ",A3)'
    ws["G6"] = "=1/0"
    ws["A7"] = True
    ws["B7"] = "=AND(B2>5,B3<5)"
    ws["C7"] = "=COUNTIF(B2:B4,\">5\")"
    ws["D7"] = "=SUMIF(A2:A4,\"Çay\",D2:D4)"
    ws["E7"] = "=IFERROR(1/0,-1)"
    ws["F7"] = "=LEN(G4)+MAX(B2:B4)-MIN(B2:B4)"
    ws["G7"] = '=UPPER(LEFT(G2,2))&MID("abc",2,1)&RIGHT("xyz",1)'
    ws["A8"] = datetime.time(10, 30)
    ws["A8"].number_format = "HH:MM"
    ws["B8"] = "=$B$2+Özet!B1"
    ws["C8"] = "=ABS(-3)+INT(2.7)+MOD(7,3)"
    ws["D8"] = "=Toplam_Adet*2"
    ws["A10"] = "Birleşik başlık"
    ws.merge_cells("A10:C11")
    ws["A10"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws["A10"].fill = PatternFill("solid", fgColor="FFFF00")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["G"].width = 24
    ws.row_dimensions[10].height = 30
    wb.defined_names["Toplam_Adet"] = DefinedName("Toplam_Adet", attr_text="'Satışlar'!$B$2:$B$4")

    oz = wb.create_sheet("Özet")
    oz["A1"] = "Satış toplamı"
    oz["B1"] = "='Satışlar'!D5"
    oz["A2"] = "Büyük harf"
    oz["B2"] = '=UPPER("istanbul")'
    oz["A3"] = 1234567.891
    oz["A3"].number_format = "#,##0.000"
    oz["A4"] = -42
    oz["A5"] = "=NOT(TRUE)"
    oz["A6"] = '=TRIM("  a  b  ")&LOWER("ÇAY")'

    uz = wb.create_sheet("Uzun")
    for i in range(1, 361):
        # Türkçe (16 bit) ve ASCII (sıkıştırılmış) dizgiler karışık: SST CONTINUE'ya bölünür
        if i % 3 == 0:
            uz.cell(i, 1, f"Satır {i} — çğıöşü ÇĞİÖŞÜ uzun metin {i * 7}")
        else:
            uz.cell(i, 1, f"Row {i} plain ascii text value {i * 13}")
        uz.cell(i, 2, i * 1.5)
        uz.cell(i, 3, f"=B{i}*2")
    gz = wb.create_sheet("Gizli")
    gz["A1"] = "gizli sayfa"
    gz.sheet_state = "hidden"
    wb.save(path)


def tarih1904(path):
    wb = Workbook()
    wb.epoch = CALENDAR_MAC_1904
    ws = wb.active
    ws.title = "Tarih"
    ws["A1"] = datetime.date(2024, 3, 15)
    ws["A1"].number_format = "DD.MM.YYYY"
    ws["A2"] = 5
    ws["A3"] = "=A1+A2"
    ws["A3"].number_format = "DD.MM.YYYY"
    wb.save(path)


def sablon(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Şablon"
    ws["A1"] = "Fatura"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = 3
    ws["B2"] = "=A2*10"
    wb.save(path)


def sifreli(src, dst, outdir):
    """Parolalı .xls (FILEPASS) — LibreOffice UNO ile kaydedilir."""
    import time
    import uno
    from com.sun.star.beans import PropertyValue

    profil = tempfile.mkdtemp(prefix="oh-lo-profil-")
    boru = f"ohtablo{os.getpid()}"
    proc = subprocess.Popen(
        ["soffice", "--headless", "--invisible", "--norestore", f"-env:UserInstallation=file://{profil}", f"--accept=pipe,name={boru};urp;"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )

    def pv(ad, deger):
        p = PropertyValue()
        p.Name = ad
        p.Value = deger
        return p

    try:
        yerel = uno.getComponentContext()
        cozucu = yerel.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", yerel)
        ctx = None
        for _ in range(120):
            try:
                ctx = cozucu.resolve(f"uno:pipe,name={boru};urp;StarOffice.ComponentContext")
                break
            except Exception:
                time.sleep(0.5)
        masa = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
        belge = masa.loadComponentFromURL(uno.systemPathToFileUrl(src), "_blank", 0, (pv("Hidden", True),))
        belge.storeToURL(uno.systemPathToFileUrl(dst), (pv("FilterName", "MS Excel 97"), pv("Password", "1234")))
        belge.close(True)
        try:
            masa.terminate()
        except Exception:
            pass
    finally:
        try:
            proc.wait(timeout=20)
        except Exception:
            proc.kill()
        shutil.rmtree(profil, ignore_errors=True)


def makrolu(src, dst, sablon_mu=False):
    """LibreOffice'in .xlsm çıktısına örnek bir makro parçası (vbaProject.bin) ekler."""
    import zipfile

    ana = "application/vnd.ms-excel.template.macroEnabled.main+xml" if sablon_mu else "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            veri = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                s = veri.decode("utf-8")
                s = s.replace("application/vnd.ms-excel.sheet.macroEnabled.main+xml", ana)
                s = s.replace("<Default Extension=\"xml\"", "<Default Extension=\"bin\" ContentType=\"application/vnd.ms-office.vbaProject\"/><Default Extension=\"xml\"", 1)
                veri = s.encode("utf-8")
            elif item.filename == "xl/_rels/workbook.xml.rels":
                s = veri.decode("utf-8")
                s = s.replace("</Relationships>", '<Relationship Id="rIdVBA" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>')
                veri = s.encode("utf-8")
            zout.writestr(item, veri)
        zout.writestr("xl/vbaProject.bin", b"OfficeHere test: makro yer tutucu")


def cevir(src, ext, outdir, filtre=None):
    hedef = f"{ext}:{filtre}" if filtre else ext
    subprocess.run(
        ["timeout", "120", "soffice", "--headless", "--convert-to", hedef, src, "--outdir", outdir],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )


def main():
    tmp = tempfile.mkdtemp(prefix="oh-fikstur-")
    try:
        ornek = os.path.join(tmp, "ornek.xlsx")
        ornek_kitap(ornek)
        cevir(ornek, "ods", tmp)
        cevir(ornek, "xls", tmp, "MS Excel 97")
        t1904 = os.path.join(tmp, "tarih1904.xlsx")
        tarih1904(t1904)
        cevir(t1904, "xls", tmp, "MS Excel 97")
        cevir(t1904, "ods", tmp)
        sb = os.path.join(tmp, "sablon.xlsx")
        sablon(sb)
        cevir(sb, "ots", tmp, "calc8_template")
        cevir(sb, "xltx", tmp, "Calc MS Excel 2007 XML Template")
        cevir(sb, "xlsm", tmp, "Calc MS Excel 2007 VBA XML")
        os.rename(os.path.join(tmp, "sablon.xlsm"), os.path.join(tmp, "lo.xlsm"))
        makrolu(os.path.join(tmp, "lo.xlsm"), os.path.join(tmp, "sablon.xlsm"))
        makrolu(os.path.join(tmp, "lo.xlsm"), os.path.join(tmp, "sablon.xltm"), sablon_mu=True)
        sifreli(sb, os.path.join(tmp, "sifreli.xls"), tmp)
        for ad in ["sifreli.xls", "sablon.xltm", "ornek.ods", "ornek.xls", "tarih1904.xls", "tarih1904.ods", "sablon.ots", "sablon.xltx", "sablon.xlsm"]:
            shutil.copy(os.path.join(tmp, ad), os.path.join(HERE, ad))
            print(ad, os.path.getsize(os.path.join(HERE, ad)), "bayt")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())

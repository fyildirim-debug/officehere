#!/usr/bin/env python3
"""Performans ölçümü için büyük dosyalar üretir (bizim içeriğimiz; tohumlu):
- buyuk-belge.docx : ~300 sayfa metin + 50 tablo
- buyuk-tablo.xlsx : 200.000 satır x 10 sütun
- zincir-10k.xlsx  : 10.000 formüllük zincir (A1=1, An=A(n-1)+1) + toplam
- buyuk-sunu.pptx  : 200 slayt (başlık, maddeler, şekil, her 10 slaytta resim)
Kurulum: pip install python-docx python-pptx xlsxwriter
Kullanım: python3 scripts/derlem/olcum/perf_veri.py [çıktı_dizini]   (varsayılan .derlem/perf/veri)
"""
import os
import random
import sys

KOK = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
CIKTI = os.environ.get("DERLEM_CIKTI", os.path.join(KOK, ".derlem"))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(CIKTI, "perf", "veri")
IMG = os.environ.get("DERLEM_GORSEL", os.path.join(KOK, "test-derlem", "gorsel"))
os.makedirs(OUT, exist_ok=True)
random.seed(7)
WORDS = "ölçüm belge tablo sunum paragraf satır sütun hücre biçim yazı tipi kalın italik başlık içerik metin sayfa düzen kenar boşluk girinti hizalama güncelleme kayıt dosya açılış kapanış Türkçe karakter ığüşöç İĞÜŞÖÇ".split()


def sentence(n=14):
    w = [random.choice(WORDS) for _ in range(n)]
    return " ".join(w).capitalize() + "."


# ---------------- DOCX ~300 sayfa
from docx import Document

d = Document()
d.add_heading("Büyük Belge — Performans Ölçümü", 0)
para_count = 0
for ch in range(1, 61):  # 60 bölüm x ~5 sayfa
    d.add_heading(f"Bölüm {ch}", 1)
    for k in range(40):
        p = d.add_paragraph(" ".join(sentence() for _ in range(4)))
        if k % 7 == 0:
            p.add_run(" Kalın ek.").bold = True
        para_count += 1
    if ch <= 50:
        t = d.add_table(rows=8, cols=5)
        t.style = "Table Grid"
        for r in range(8):
            for c in range(5):
                t.cell(r, c).text = f"T{ch}-{r}{c}"
d.add_paragraph("SON-PARAGRAF-İŞARETİ")
d.save(os.path.join(OUT, "buyuk-belge.docx"))
print("docx paragraf", para_count)

# ---------------- XLSX 200k x 10 (xlsxwriter, sabit bellek)
import xlsxwriter

wb = xlsxwriter.Workbook(os.path.join(OUT, "buyuk-tablo.xlsx"), {"constant_memory": True})
ws = wb.add_worksheet("Veri")
cities = ["İstanbul", "Ankara", "İzmir", "Bursa", "Antalya", "Konya", "Çanakkale"]
ws.write_row(0, 0, ["No", "Ad", "Şehir", "Adet", "Fiyat", "Tutar", "Tarih", "Oran", "Kod", "Not"])
for r in range(1, 200001):
    ws.write_row(r, 0, [r, f"Müşteri {r}", cities[r % 7], (r * 7) % 50, 10 + (r % 13), ((r * 7) % 50) * (10 + (r % 13)), 45000 + (r % 900), (r % 100) / 100, f"K{r:06d}", "not" if r % 5 == 0 else ""])
wb.close()
print("xlsx 200k tamam")

# ---------------- XLSX formül zinciri 10k
wb = xlsxwriter.Workbook(os.path.join(OUT, "zincir-10k.xlsx"))
ws = wb.add_worksheet("Zincir")
ws.write_number(0, 0, 1)
for r in range(1, 10000):
    ws.write_formula(r, 0, f"=A{r}+1", None, r + 1)
ws.write_formula(0, 1, "=SUM(A1:A10000)", None, 50005000)
ws.write_formula(1, 1, "=A10000*2", None, 20000)
wb.close()
print("zincir tamam")

# ---------------- PPTX 200 slayt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

pr = Presentation()
pr.slide_width, pr.slide_height = Inches(13.333), Inches(7.5)
img = os.path.join(IMG, "foto.jpg")
for i in range(200):
    s = pr.slides.add_slide(pr.slide_layouts[1])
    s.shapes.title.text = f"Slayt {i + 1}: {sentence(5)}"
    tf = s.placeholders[1].text_frame
    tf.text = sentence(8)
    for j in range(4):
        p = tf.add_paragraph()
        p.text = sentence(7)
        p.level = j % 2
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9), Inches(5.5), Inches(3.5), Inches(1.4))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(0x44, 0x72, 0xC4)
    sh.text = f"Şekil {i + 1}"
    if i % 10 == 0:
        s.shapes.add_picture(img, Inches(9), Inches(1.8), width=Inches(3.5))
    s.notes_slide.notes_text_frame.text = f"Not {i + 1}: " + sentence(20)
pr.save(os.path.join(OUT, "buyuk-sunu.pptx"))
print("pptx tamam")
for f in sorted(os.listdir(OUT)):
    print(f, os.path.getsize(os.path.join(OUT, f)))

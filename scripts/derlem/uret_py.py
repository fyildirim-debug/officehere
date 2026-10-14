#!/usr/bin/env python3
"""İkinci üretici (Office benzeri OOXML): python-docx, openpyxl ve python-pptx ile derlem üretir.

LibreOffice'in yazdığından farklı kalıpları (varsayılan şablon stilleri, grafik parçaları,
koşullu biçim kuralları…) sınar. İçerik bizimdir; kütüphaneler yalnız üretimde kullanılır
(dağıtılmaz). Kurulum: pip install python-docx openpyxl python-pptx

Kullanım: python3 scripts/derlem/uret_py.py [çıktı_dizini]   (varsayılan .derlem/uretilen/py)
Görseller: DERLEM_GORSEL (varsayılan test-derlem/gorsel).
"""
import os
import sys

KOK = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(KOK, ".derlem", "uretilen", "py")
IMG = os.environ.get("DERLEM_GORSEL", os.path.join(KOK, "test-derlem", "gorsel"))
os.makedirs(OUT, exist_ok=True)
TXT = "Python ile üretilmiş paragraf; Türkçe karakterler ığüşöç İĞÜŞÖÇ ve uzun cümleler içerir. " * 3

# ---------------- DOCX
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt, RGBColor

d = Document()
d.core_properties.title = "python-docx derlem"
d.core_properties.author = "Ölçüm"
sec = d.sections[0]
sec.header.paragraphs[0].text = "Üst bilgi: python-docx"
sec.footer.paragraphs[0].text = "Alt bilgi metni"
d.add_heading("python-docx Derlem Belgesi", 0)
for s in range(1, 13):
    d.add_heading(f"Bölüm {s}", 1)
    p = d.add_paragraph(TXT)
    r = p.add_run(" kalın")
    r.bold = True
    r = p.add_run(" italik")
    r.italic = True
    r = p.add_run(" renkli")
    r.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
    r.font.size = Pt(15)
    r.font.name = "Courier New"
    p2 = d.add_paragraph(TXT)
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i in range(3):
        d.add_paragraph(f"Madde {i + 1}", style="List Bullet")
    for i in range(3):
        d.add_paragraph(f"Numaralı {i + 1}", style="List Number")
    d.add_paragraph("İç madde", style="List Bullet 2")
    if s % 3 == 0:
        t = d.add_table(rows=6, cols=5)
        t.style = "Table Grid"
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                cell.text = f"Başlık {ci}" if ri == 0 else f"{ri}-{ci}"
        a = t.cell(5, 0).merge(t.cell(5, 2))
        a.text = "birleşik"
    if s % 4 == 0:
        d.add_picture(os.path.join(IMG, "grafik.png" if s % 8 else "foto.jpg"), width=Inches(3))
    if s % 5 == 0:
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
ns = d.add_section()
ns.orientation = WD_ORIENT.LANDSCAPE
ns.page_width, ns.page_height = ns.page_height, ns.page_width
d.add_paragraph("Yatay bölüm metni. " + TXT)
d.save(os.path.join(OUT, "py-belge.docx"))

# ---------------- XLSX
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
ws = wb.active
ws.title = "Veri"
ws.append(["Ürün", "Adet", "Fiyat", "Tutar", "Tarih"])
import datetime as dt

for i in range(1, 101):
    ws.append([f"Ürün {i} ğüş", i % 17 + 1, round(3.5 + i * 0.25, 2), f"=B{i + 1}*C{i + 1}", dt.date(2024, (i % 12) + 1, (i % 27) + 1)])
ws["G1"] = "Toplam"
ws["H1"] = "=SUM(D2:D101)"
ws["G2"] = "Ortalama"
ws["H2"] = "=AVERAGE(D2:D101)"
ws["G3"] = "Eğer"
ws["H3"] = '=IF(H1>1000,"Çok","Az")'
ws["G4"] = "Düşey ara"
ws["H4"] = '=VLOOKUP("Ürün 5 ğüş",A2:D101,4,FALSE)'
hdr = Font(bold=True, color="FFFFFF")
fill = PatternFill("solid", fgColor="305496")
thin = Side(style="thin", color="000000")
for c in ws[1]:
    c.font = hdr
    c.fill = fill
    c.alignment = Alignment(horizontal="center")
    c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for r in range(2, 102):
    ws.cell(r, 3).number_format = "#,##0.00"
    ws.cell(r, 5).number_format = "dd.mm.yyyy"
ws.merge_cells("J1:L2")
ws["J1"] = "Birleşik alan"
ws.freeze_panes = "B2"
ws.column_dimensions["A"].width = 22
ws["A1"].comment = Comment("Ürün adları", "Ölçüm")
dv = DataValidation(type="list", formula1='"Evet,Hayır"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("M2:M20")
ws.conditional_formatting.add("D2:D101", CellIsRule(operator="greaterThan", formula=["100"], fill=PatternFill("solid", fgColor="FFC7CE")))
ws.conditional_formatting.add("B2:B101", ColorScaleRule(start_type="min", start_color="F8696B", end_type="max", end_color="63BE7B"))
ch = BarChart()
ch.title = "Adetler"
ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=15), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=15))
ws.add_chart(ch, "O2")
lc = LineChart()
lc.add_data(Reference(ws, min_col=3, min_row=1, max_row=30), titles_from_data=True)
ws.add_chart(lc, "O20")
ws2 = wb.create_sheet("Özet")
ws2["A1"] = "Toplam"
ws2["B1"] = "=Veri!H1"
ws2["A2"] = "Sayım"
ws2["B2"] = "=COUNTA(Veri!A2:A101)"
ws2["A3"] = "Metin"
ws2["B3"] = '=CONCATENATE("Toplam=",TEXT(B1,"0.00"))'
pc = PieChart()
pc.add_data(Reference(ws, min_col=2, min_row=1, max_row=6), titles_from_data=True)
ws2.add_chart(pc, "D2")
wb.save(os.path.join(OUT, "py-tablo.xlsx"))

# ---------------- PPTX
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches as PI, Pt as PPt

pr = Presentation()
s = pr.slides.add_slide(pr.slide_layouts[0])
s.shapes.title.text = "python-pptx Derlem"
s.placeholders[1].text = "Alt başlık ığüşöç"
for i in range(12):
    s = pr.slides.add_slide(pr.slide_layouts[1])
    s.shapes.title.text = f"İçerik slaytı {i + 1}"
    tf = s.placeholders[1].text_frame
    tf.text = "Birinci madde"
    p = tf.add_paragraph()
    p.text = "İç madde"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Kalın madde"
    p.runs[0].font.bold = True
    s.notes_slide.notes_text_frame.text = f"Not {i + 1}: anlatım"
    if i % 4 == 1:
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PI(6), PI(4.5), PI(3), PI(1.5))
        sh.text = "Yuvarlak dikdörtgen"
        sh.fill.solid()
        sh.fill.fore_color.rgb = PRGB(0xED, 0x7D, 0x31)
        st = s.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, PI(1), PI(5), PI(1.5), PI(1.5))
    if i % 4 == 2:
        rows, cols = 4, 3
        tbl = s.shapes.add_table(rows, cols, PI(1), PI(4.2), PI(8), PI(2)).table
        for r in range(rows):
            for c in range(cols):
                tbl.cell(r, c).text = f"H{r}{c}"
        tbl.cell(3, 0).merge(tbl.cell(3, 1))
    if i % 4 == 3:
        s.shapes.add_picture(os.path.join(IMG, "foto.jpg"), PI(5.5), PI(3.5), width=PI(3.5))
    if i % 4 == 0:
        cd = CategoryChartData()
        cd.categories = ["Oca", "Şub", "Mar"]
        cd.add_series("Satış", (12.5, 18.0, 9.75))
        cd.add_series("Hedef", (15, 15, 15))
        s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, PI(5), PI(4), PI(4.5), PI(3), cd)
pr.save(os.path.join(OUT, "py-sunu.pptx"))
print("tamam", os.listdir(OUT))
